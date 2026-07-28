"""
Excel / CSV import and export for student and faculty records.

Import is deliberately **strict and transactional per row**: every row is
validated before any write, and a row that fails is reported with its line
number and the exact reason rather than being silently skipped.  The caller
receives an :class:`ImportResult` it can render as a report.

A blank template with data-validation dropdowns can be generated so the office
staff enter branch/semester names that actually exist.
"""

from __future__ import annotations

from datetime import datetime
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from config.settings import EXPORT_DIR, IMPORT_DIR, config
from core.audit import ACTION_EXPORT, ACTION_IMPORT, log_audit
from core.database import get_db
from core.logger import get_logger
from core.validators import (parse_date, validate_email, validate_enrollment,
                             validate_mobile, validate_name, normalise_mobile)
from models import academic, faculty as faculty_model, student as student_model

logger = get_logger("services.import_export")

# Column headings expected in a student import file.
STUDENT_COLUMNS = [
    "Enrollment No", "Roll No", "Full Name", "Father Name", "Mother Name",
    "Gender", "DOB (YYYY-MM-DD)", "Mobile", "Email", "Address",
    "Admission Date (YYYY-MM-DD)", "Branch Code", "Semester Number",
    "Section", "Batch", "Status",
]
STUDENT_REQUIRED = ["Enrollment No", "Roll No", "Full Name", "Branch Code",
                    "Semester Number"]

FACULTY_COLUMNS = [
    "Faculty Code", "Full Name", "Gender", "DOB (YYYY-MM-DD)", "Qualification",
    "Designation", "Experience (Years)", "Mobile", "Email", "Address",
    "Branch Code", "Joining Date (YYYY-MM-DD)", "Status",
]
FACULTY_REQUIRED = ["Faculty Code", "Full Name"]


class ImportResult:
    """Outcome of an import run, ready to be shown to the user."""

    def __init__(self) -> None:
        self.imported = 0
        self.skipped = 0
        self.errors: list[tuple[int, str, str]] = []   # (row_no, identifier, reason)
        self.warnings: list[str] = []

    @property
    def total(self) -> int:
        return self.imported + self.skipped

    @property
    def success(self) -> bool:
        return self.imported > 0

    def add_error(self, row_no: int, identifier: str, reason: str) -> None:
        self.errors.append((row_no, identifier, reason))
        self.skipped += 1

    def summary(self) -> str:
        lines = [f"Imported: {self.imported}", f"Skipped:  {self.skipped}"]
        if self.errors:
            lines.append("")
            lines.append("Problems found:")
            for row_no, identifier, reason in self.errors[:40]:
                lines.append(f"  Row {row_no} ({identifier or '-'}): {reason}")
            if len(self.errors) > 40:
                lines.append(f"  ... and {len(self.errors) - 40} more.")
        if self.warnings:
            lines.append("")
            lines.extend(f"Note: {w}" for w in self.warnings)
        return "\n".join(lines)


# ---------------------------------------------------------------------------
# Lookup caches -- resolve human-typed codes to primary keys
# ---------------------------------------------------------------------------
def _lookups() -> dict:
    db = get_db()
    return {
        "branches": {str(r["branch_code"]).upper(): r["branch_id"]
                     for r in db.fetch_all("SELECT branch_id, branch_code FROM branches")},
        "branch_courses": {r["branch_id"]: r["course_id"]
                           for r in db.fetch_all("SELECT branch_id, course_id FROM branches")},
        "semesters": {(r["course_id"], int(r["semester_number"])): r["semester_id"]
                      for r in db.fetch_all(
                          "SELECT semester_id, course_id, semester_number FROM semesters")},
        "sections": {str(r["section_name"]).upper(): r["section_id"]
                     for r in db.fetch_all("SELECT section_id, section_name FROM sections")},
        "batches": {str(r["batch_name"]): r["batch_id"]
                    for r in db.fetch_all("SELECT batch_id, batch_name FROM batches")},
    }


def _read_table(path: str | Path) -> pd.DataFrame:
    """Read XLSX or CSV into a DataFrame with clean string cells."""
    path = Path(path)
    if path.suffix.lower() in (".xlsx", ".xlsm", ".xls"):
        frame = pd.read_excel(path, dtype=str)
    elif path.suffix.lower() == ".csv":
        frame = pd.read_csv(path, dtype=str, comment="#")
    else:
        raise ValueError(f"Unsupported file type '{path.suffix}'. Use .xlsx or .csv.")

    frame.columns = [str(c).strip() for c in frame.columns]
    # Normalise every cell to a trimmed string; NaN becomes empty.
    return frame.fillna("").apply(lambda col: col.map(lambda v: str(v).strip()))


def _cell(row, column: str) -> str:
    value = row.get(column, "")
    return "" if value in ("nan", "NaT", "None") else str(value).strip()


def _iso_date(value: str) -> str | None:
    parsed = parse_date(value)
    return parsed.isoformat() if parsed else None


# ===========================================================================
# Student import
# ===========================================================================
def import_students(path: str | Path, session_id: int | None = None,
                    default_status: str = "Active",
                    user: dict | None = None) -> ImportResult:
    """Import students from an Excel/CSV file."""
    result = ImportResult()

    try:
        frame = _read_table(path)
    except Exception as exc:                    # noqa: BLE001
        result.add_error(0, "", f"Could not read the file: {exc}")
        return result

    missing = [c for c in STUDENT_REQUIRED if c not in frame.columns]
    if missing:
        result.add_error(0, "", f"Missing required column(s): {', '.join(missing)}")
        return result

    lookups = _lookups()
    session_id = session_id or academic.get_current_session_id()
    db = get_db()

    seen_enrollments: set[str] = set()

    for index, row in frame.iterrows():
        row_no = int(index) + 2           # +2: header row plus 1-based indexing
        enrollment = _cell(row, "Enrollment No")

        if not any(_cell(row, c) for c in STUDENT_REQUIRED):
            continue                      # entirely blank row

        # -- identity ------------------------------------------------------
        ok, message = validate_enrollment(enrollment)
        if not ok:
            result.add_error(row_no, enrollment, message)
            continue
        if enrollment.upper() in seen_enrollments:
            result.add_error(row_no, enrollment, "Duplicate enrollment number in this file.")
            continue
        if db.exists("students", "enrollment_no = ?", (enrollment,)):
            result.add_error(row_no, enrollment, "Enrollment number already exists in the database.")
            continue
        seen_enrollments.add(enrollment.upper())

        full_name = _cell(row, "Full Name")
        ok, message = validate_name(full_name, "Full Name")
        if not ok:
            result.add_error(row_no, enrollment, message)
            continue

        roll_no = _cell(row, "Roll No")
        if not roll_no:
            result.add_error(row_no, enrollment, "Roll No is required.")
            continue

        # -- academic mapping ---------------------------------------------
        branch_code = _cell(row, "Branch Code").upper()
        branch_id = lookups["branches"].get(branch_code)
        if not branch_id:
            result.add_error(row_no, enrollment,
                             f"Unknown branch code '{branch_code}'. "
                             f"Known: {', '.join(sorted(lookups['branches'])) or 'none'}")
            continue

        course_id = lookups["branch_courses"].get(branch_id)
        try:
            semester_number = int(float(_cell(row, "Semester Number")))
        except (TypeError, ValueError):
            result.add_error(row_no, enrollment, "Semester Number must be a whole number.")
            continue

        semester_id = lookups["semesters"].get((course_id, semester_number))
        if not semester_id:
            result.add_error(row_no, enrollment,
                             f"Semester {semester_number} does not exist for this course.")
            continue

        section_name = _cell(row, "Section").upper()
        section_id = lookups["sections"].get(section_name) if section_name else None
        if section_name and not section_id:
            result.warnings.append(
                f"Row {row_no}: section '{section_name}' does not exist; left blank.")

        batch_name = _cell(row, "Batch")
        batch_id = lookups["batches"].get(batch_name) if batch_name else None
        if batch_name and not batch_id:
            result.warnings.append(
                f"Row {row_no}: batch '{batch_name}' does not exist; left blank.")

        # -- optional contact details -------------------------------------
        mobile = normalise_mobile(_cell(row, "Mobile"))
        if mobile:
            ok, message = validate_mobile(mobile)
            if not ok:
                result.add_error(row_no, enrollment, message)
                continue

        email = _cell(row, "Email")
        if email:
            ok, message = validate_email(email)
            if not ok:
                result.add_error(row_no, enrollment, message)
                continue

        status = _cell(row, "Status") or default_status
        if status not in ("Active", "Inactive", "Alumni", "Dropped"):
            status = default_status

        payload = {
            "enrollment_no": enrollment, "roll_no": roll_no, "full_name": full_name,
            "father_name": _cell(row, "Father Name"),
            "mother_name": _cell(row, "Mother Name"),
            "gender": _cell(row, "Gender") or None,
            "dob": _iso_date(_cell(row, "DOB (YYYY-MM-DD)")),
            "mobile": mobile or None, "email": email or None,
            "address": _cell(row, "Address"),
            "admission_date": _iso_date(_cell(row, "Admission Date (YYYY-MM-DD)")),
            "course_id": course_id, "branch_id": branch_id, "semester_id": semester_id,
            "section_id": section_id, "batch_id": batch_id, "session_id": session_id,
            "status": status,
        }

        ok, message, _ = student_model.add_student(payload, user=user)
        if ok:
            result.imported += 1
        else:
            result.add_error(row_no, enrollment, message)

    log_audit(user, ACTION_IMPORT, "Students", "students", None,
              new_value={"file": Path(path).name, "imported": result.imported,
                         "skipped": result.skipped})
    logger.info("Student import from %s: %d imported, %d skipped",
                Path(path).name, result.imported, result.skipped)
    return result


# ===========================================================================
# Faculty import
# ===========================================================================
def import_faculty(path: str | Path, user: dict | None = None) -> ImportResult:
    result = ImportResult()

    try:
        frame = _read_table(path)
    except Exception as exc:                    # noqa: BLE001
        result.add_error(0, "", f"Could not read the file: {exc}")
        return result

    missing = [c for c in FACULTY_REQUIRED if c not in frame.columns]
    if missing:
        result.add_error(0, "", f"Missing required column(s): {', '.join(missing)}")
        return result

    lookups = _lookups()
    db = get_db()
    seen: set[str] = set()

    for index, row in frame.iterrows():
        row_no = int(index) + 2
        code = _cell(row, "Faculty Code")

        if not any(_cell(row, c) for c in FACULTY_REQUIRED):
            continue

        if not code:
            result.add_error(row_no, "", "Faculty Code is required.")
            continue
        if code.upper() in seen:
            result.add_error(row_no, code, "Duplicate faculty code in this file.")
            continue
        if db.exists("faculty", "faculty_code = ?", (code,)):
            result.add_error(row_no, code, "Faculty code already exists in the database.")
            continue
        seen.add(code.upper())

        full_name = _cell(row, "Full Name")
        ok, message = validate_name(full_name, "Full Name")
        if not ok:
            result.add_error(row_no, code, message)
            continue

        branch_code = _cell(row, "Branch Code").upper()
        branch_id = lookups["branches"].get(branch_code) if branch_code else None
        if branch_code and not branch_id:
            result.warnings.append(
                f"Row {row_no}: branch '{branch_code}' not found; left unassigned.")

        mobile = normalise_mobile(_cell(row, "Mobile"))
        if mobile:
            ok, message = validate_mobile(mobile)
            if not ok:
                result.add_error(row_no, code, message)
                continue

        try:
            experience = float(_cell(row, "Experience (Years)") or 0)
        except ValueError:
            experience = 0.0

        payload = {
            "faculty_code": code, "full_name": full_name,
            "gender": _cell(row, "Gender") or None,
            "dob": _iso_date(_cell(row, "DOB (YYYY-MM-DD)")),
            "qualification": _cell(row, "Qualification"),
            "designation": _cell(row, "Designation"),
            "experience_years": experience,
            "mobile": mobile or None, "email": _cell(row, "Email") or None,
            "address": _cell(row, "Address"), "branch_id": branch_id,
            "joining_date": _iso_date(_cell(row, "Joining Date (YYYY-MM-DD)")),
            "status": _cell(row, "Status") or "Active",
        }

        ok, message, _ = faculty_model.add_faculty(payload, user=user)
        if ok:
            result.imported += 1
        else:
            result.add_error(row_no, code, message)

    log_audit(user, ACTION_IMPORT, "Faculty", "faculty", None,
              new_value={"file": Path(path).name, "imported": result.imported,
                         "skipped": result.skipped})
    return result


# ===========================================================================
# Export
# ===========================================================================
def _write_sheet(rows: list[dict], columns: list[str], title: str,
                 path: Path) -> None:
    """Shared writer for the export helpers, with the college header band."""
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = title[:31]

    last_column = get_column_letter(max(len(columns), 3))
    sheet.merge_cells(f"A1:{last_column}1")
    sheet["A1"] = config.get("college_name", "")
    sheet["A1"].font = Font(size=14, bold=True, color="0F2A4A")
    sheet["A1"].alignment = Alignment(horizontal="center")

    sheet.merge_cells(f"A2:{last_column}2")
    sheet["A2"] = f"{title}  -  exported {datetime.now():%d %b %Y %I:%M %p}"
    sheet["A2"].font = Font(size=9, color="64748B")
    sheet["A2"].alignment = Alignment(horizontal="center")

    header_row = 4
    for index, name in enumerate(columns, start=1):
        cell = sheet.cell(row=header_row, column=index, value=name)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", start_color="0F2A4A")
        cell.alignment = Alignment(horizontal="center", wrap_text=True)

    for offset, record in enumerate(rows, start=1):
        for index, name in enumerate(columns, start=1):
            sheet.cell(row=header_row + offset, column=index, value=record.get(name, ""))

    for index, name in enumerate(columns, start=1):
        longest = max([len(str(name))] +
                      [len(str(r.get(name, ""))) for r in rows[:200]] or [12])
        sheet.column_dimensions[get_column_letter(index)].width = min(max(longest + 3, 12), 40)

    sheet.freeze_panes = sheet.cell(row=header_row + 1, column=1)
    workbook.save(path)


def export_students(students: list, file_format: str = "Excel",
                    user: dict | None = None) -> tuple[bool, str, Path | None]:
    """Export student rows to Excel or CSV."""
    EXPORT_DIR.mkdir(parents=True, exist_ok=True)
    stamp = datetime.now().strftime("%Y%m%d_%H%M%S")

    records = [{
        "Enrollment No": s["enrollment_no"], "Roll No": s["roll_no"],
        "Full Name": s["full_name"], "Father Name": s["father_name"] or "",
        "Mother Name": s["mother_name"] or "", "Gender": s["gender"] or "",
        "DOB (YYYY-MM-DD)": s["dob"] or "", "Mobile": s["mobile"] or "",
        "Email": s["email"] or "", "Address": s["address"] or "",
        "Admission Date (YYYY-MM-DD)": s["admission_date"] or "",
        "Branch Code": s["branch_code"], "Semester Number": s["semester_number"],
        "Section": s["section_name"] or "", "Batch": s["batch_name"] or "",
        "Status": s["status"],
    } for s in students]

    try:
        if file_format.upper() in ("EXCEL", "XLSX"):
            path = EXPORT_DIR / f"Students_Export_{stamp}.xlsx"
            _write_sheet(records, STUDENT_COLUMNS, "Student Records", path)
        else:
            path = EXPORT_DIR / f"Students_Export_{stamp}.csv"
            pd.DataFrame(records, columns=STUDENT_COLUMNS).to_csv(
                path, index=False, encoding="utf-8-sig")

        log_audit(user, ACTION_EXPORT, "Students", "students", None,
                  new_value={"count": len(records), "file": path.name})
        return True, f"{len(records)} student(s) exported to {path.name}", path

    except Exception as exc:                    # noqa: BLE001
        logger.error("Student export failed: %s", exc, exc_info=True)
        return False, f"Export failed: {exc}", None


def export_faculty(faculty_rows: list, file_format: str = "Excel",
                   user: dict | None = None) -> tuple[bool, str, Path | None]:
    EXPORT_DIR.mkdir(parents=True, exist_ok=True)
    stamp = datetime.now().strftime("%Y%m%d_%H%M%S")

    records = [{
        "Faculty Code": f["faculty_code"], "Full Name": f["full_name"],
        "Gender": f["gender"] or "", "DOB (YYYY-MM-DD)": f["dob"] or "",
        "Qualification": f["qualification"] or "", "Designation": f["designation"] or "",
        "Experience (Years)": f["experience_years"] or 0,
        "Mobile": f["mobile"] or "", "Email": f["email"] or "",
        "Address": f["address"] or "",
        "Branch Code": f["branch_code"] if "branch_code" in f.keys() else "",
        "Joining Date (YYYY-MM-DD)": f["joining_date"] or "", "Status": f["status"],
    } for f in faculty_rows]

    try:
        if file_format.upper() in ("EXCEL", "XLSX"):
            path = EXPORT_DIR / f"Faculty_Export_{stamp}.xlsx"
            _write_sheet(records, FACULTY_COLUMNS, "Faculty Records", path)
        else:
            path = EXPORT_DIR / f"Faculty_Export_{stamp}.csv"
            pd.DataFrame(records, columns=FACULTY_COLUMNS).to_csv(
                path, index=False, encoding="utf-8-sig")

        log_audit(user, ACTION_EXPORT, "Faculty", "faculty", None,
                  new_value={"count": len(records), "file": path.name})
        return True, f"{len(records)} faculty record(s) exported to {path.name}", path

    except Exception as exc:                    # noqa: BLE001
        logger.error("Faculty export failed: %s", exc, exc_info=True)
        return False, f"Export failed: {exc}", None


# ===========================================================================
# Templates
# ===========================================================================
def create_student_template(path: Path | None = None) -> tuple[bool, str, Path | None]:
    """Blank student import workbook with dropdowns and an instructions sheet."""
    IMPORT_DIR.mkdir(parents=True, exist_ok=True)
    path = path or (IMPORT_DIR / "Student_Import_Template.xlsx")

    try:
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Students"

        for index, name in enumerate(STUDENT_COLUMNS, start=1):
            cell = sheet.cell(row=1, column=index, value=name)
            required = name in STUDENT_REQUIRED
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid",
                                    start_color="B91C1C" if required else "0F2A4A")
            cell.alignment = Alignment(horizontal="center", wrap_text=True)
            sheet.column_dimensions[get_column_letter(index)].width = max(len(name) + 4, 15)
        sheet.row_dimensions[1].height = 30

        # -- reference sheet powers the dropdowns --------------------------
        reference = workbook.create_sheet("Reference")
        branches = get_db().fetch_all(
            "SELECT branch_code, branch_name FROM branches WHERE is_active = 1")
        sections = get_db().fetch_all("SELECT section_name FROM sections WHERE is_active = 1")
        batches = get_db().fetch_all("SELECT batch_name FROM batches WHERE is_active = 1")

        reference["A1"] = "Branch Code"
        reference["B1"] = "Branch Name"
        for index, row in enumerate(branches, start=2):
            reference[f"A{index}"] = row["branch_code"]
            reference[f"B{index}"] = row["branch_name"]

        reference["D1"] = "Sections"
        for index, row in enumerate(sections, start=2):
            reference[f"D{index}"] = row["section_name"]

        reference["F1"] = "Batches"
        for index, row in enumerate(batches, start=2):
            reference[f"F{index}"] = row["batch_name"]

        for column in ("A1", "B1", "D1", "F1"):
            reference[column].font = Font(bold=True, color="FFFFFF")
            reference[column].fill = PatternFill("solid", start_color="0F2A4A")
            reference.column_dimensions[column[0]].width = 26

        # -- data validation ------------------------------------------------
        def _add_validation(column_letter: str, formula: str) -> None:
            validation = DataValidation(type="list", formula1=formula, allow_blank=True)
            sheet.add_data_validation(validation)
            validation.add(f"{column_letter}2:{column_letter}500")

        if branches:
            _add_validation("L", f"=Reference!$A$2:$A${len(branches) + 1}")
        if sections:
            _add_validation("N", f"=Reference!$D$2:$D${len(sections) + 1}")
        if batches:
            _add_validation("O", f"=Reference!$F$2:$F${len(batches) + 1}")
        _add_validation("F", '"Male,Female,Other"')
        _add_validation("P", '"Active,Inactive,Alumni,Dropped"')

        semester_count = int(get_db().fetch_value(
            "SELECT MAX(semester_number) FROM semesters", (), 6) or 6)
        _add_validation("M", f'"{",".join(str(n) for n in range(1, semester_count + 1))}"')

        # -- instructions ---------------------------------------------------
        guide = workbook.create_sheet("Instructions")
        guide.column_dimensions["A"].width = 100
        lines = [
            ("HOW TO USE THIS TEMPLATE", True),
            ("", False),
            ("1. Fill in one student per row on the 'Students' sheet.", False),
            ("2. Columns with a RED heading are mandatory:", False),
            (f"   {', '.join(STUDENT_REQUIRED)}", False),
            ("3. Branch Code, Section and Batch must already exist in the system.", False),
            ("   The 'Reference' sheet lists every valid value, and those columns", False),
            ("   have dropdowns to prevent typing mistakes.", False),
            ("4. Dates must be written as YYYY-MM-DD, e.g. 2008-07-14.", False),
            ("5. Mobile numbers must be 10 digits starting with 6, 7, 8 or 9.", False),
            ("6. Enrollment numbers must be unique across the whole college.", False),
            ("7. Roll numbers must be unique within a branch + semester + section.", False),
            ("8. Do not rename, reorder or delete the heading row.", False),
            ("", False),
            ("After importing, capture each student's face dataset from", False),
            ("Students -> select a student -> Capture Face Dataset.", False),
        ]
        for index, (text, bold) in enumerate(lines, start=1):
            cell = guide.cell(row=index, column=1, value=text)
            cell.font = Font(bold=bold, size=12 if bold else 10,
                             color="0F2A4A" if bold else "000000")
            cell.alignment = Alignment(wrap_text=True)

        workbook.save(path)
        return True, f"Template created: {path}", path

    except Exception as exc:                    # noqa: BLE001
        logger.error("Template creation failed: %s", exc, exc_info=True)
        return False, f"Could not create the template: {exc}", None


def create_faculty_template(path: Path | None = None) -> tuple[bool, str, Path | None]:
    IMPORT_DIR.mkdir(parents=True, exist_ok=True)
    path = path or (IMPORT_DIR / "Faculty_Import_Template.xlsx")

    try:
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Faculty"

        for index, name in enumerate(FACULTY_COLUMNS, start=1):
            cell = sheet.cell(row=1, column=index, value=name)
            required = name in FACULTY_REQUIRED
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid",
                                    start_color="B91C1C" if required else "0F2A4A")
            cell.alignment = Alignment(horizontal="center", wrap_text=True)
            sheet.column_dimensions[get_column_letter(index)].width = max(len(name) + 4, 15)

        branches = get_db().fetch_all(
            "SELECT branch_code FROM branches WHERE is_active = 1")
        reference = workbook.create_sheet("Reference")
        reference["A1"] = "Branch Code"
        for index, row in enumerate(branches, start=2):
            reference[f"A{index}"] = row["branch_code"]

        if branches:
            validation = DataValidation(
                type="list", formula1=f"=Reference!$A$2:$A${len(branches) + 1}",
                allow_blank=True)
            sheet.add_data_validation(validation)
            validation.add("K2:K500")

        workbook.save(path)
        return True, f"Template created: {path}", path

    except Exception as exc:                    # noqa: BLE001
        return False, f"Could not create the template: {exc}", None
